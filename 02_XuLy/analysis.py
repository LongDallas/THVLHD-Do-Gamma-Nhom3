"""Phân tích số liệu bài thực tập "Phân bố thống kê trong phân rã phóng xạ".

Đọc số liệu thô từ `nhóm 1.xlsx`, thực hiện các bước của mục 2.3 và mục 3 (Poisson.pdf),
xuất đồ thị, kết quả số (results.json) và số liệu mô phỏng vào 03_KetQuaTrungGian/.

Chạy:  python3 02_XuLy/analysis.py
"""
import json
import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from scipy import stats

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "01_SoLieuGoc" / "nhóm 1.xlsx"
OUT_DIR = ROOT / "03_KetQuaTrungGian"
FIG = OUT_DIR / "figures"
MC_DIR = OUT_DIR / "montecarlo"
T_SHORT = 10.0   # thời gian mỗi lần đo (s)
T_LONG = 100.0   # thời gian phép đo dài (s)
N_MC = 1000      # số lần thử Monte Carlo
N_REP = 2000     # số mẫu MC cỡ 100 để khảo sát phân tán của s
SEED = 2026

# (khóa, nhãn, cột số liệu 10 s, hàng đầu, hàng cuối, ô phép đo 100 s)
DATASETS = [
    ("S1", "Bộ 1 – Mức thấp (Phông)", "G", 7, 106, "H7"),
    ("S2", "Bộ 2 – Mức trung bình (Ra-226)", "C", 6, 105, "C108"),
    ("S3", "Bộ 3 – Mức cao (Ra-226)", "A", 6, 105, "A108"),
]

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "figure.dpi": 150})


def load():
    ws = openpyxl.load_workbook(RAW, data_only=True).active
    data = {}
    for key, label, col, r0, r1, long_cell in DATASETS:
        x = np.array([ws[f"{col}{r}"].value for r in range(r0, r1 + 1)], dtype=int)
        data[key] = {"label": label, "x": x, "N100": int(ws[long_cell].value),
                     "src": f"{col}{r0}:{col}{r1}", "src100": long_cell}
    return data


def chi2_eq5(obs, exp):
    """χ² theo công thức (5) của tài liệu: σ_i = √x_i, bỏ các bin có x_i = 0."""
    m = obs > 0
    return float(np.sum((obs[m] - exp[m]) ** 2 / obs[m])), int(m.sum())


def chi2_pearson(x, cdf, emin=5.0):
    """χ² Pearson Σ(O−E)²/E sau khi gộp bin để E ≥ emin (kể cả hai đuôi); ndf = số nhóm − 2."""
    n = len(x)
    edges, lo = [], int(x.min())
    # Ranh giới nhóm: gộp các giá trị nguyên liên tiếp đến khi E ≥ emin
    k, prev = lo, -np.inf  # nhóm đầu gồm cả đuôi trái (−∞, k]
    while k <= x.max():
        if n * (cdf(k) - (cdf(prev) if np.isfinite(prev) else 0)) >= emin:
            edges.append(k)
            prev = k
        k += 1
    # Đuôi phải (> edges[-1]) phải đủ E ≥ emin, nếu không gộp vào nhóm trước
    while edges and n * (1 - cdf(edges[-1])) < emin:
        edges.pop()
    bounds = [-np.inf] + edges + [np.inf]
    O, E = [], []
    for a, b in zip(bounds[:-1], bounds[1:]):
        O.append(int(((x > a) & (x <= b)).sum()))
        E.append(n * ((cdf(b) if np.isfinite(b) else 1) - (cdf(a) if np.isfinite(a) else 0)))
    O, E = np.array(O), np.array(E)
    chi2 = float(((O - E) ** 2 / E).sum())
    ndf = len(O) - 2
    return {"chi2": chi2, "ndf": ndf, "p": float(stats.chi2.sf(chi2, ndf)), "groups": len(O)}


def analyse(d, rng):
    x = d["x"]
    n = len(x)
    mean = x.mean()
    s = x.std(ddof=1)
    res = {
        "label": d["label"], "src": d["src"], "src100": d["src100"], "n": n,
        "sum": int(x.sum()), "mean": mean, "s": s, "var": s ** 2,
        "u_mean": s / math.sqrt(n), "sqrt_mean": math.sqrt(mean),
        "min": int(x.min()), "max": int(x.max()),
        "rate_cps": mean / T_SHORT, "u_rate_cps": s / math.sqrt(n) / T_SHORT,
    }
    # Chỉ số phân tán D = s²/x̄ và kiểm định χ²(n−1) = (n−1)s²/x̄
    disp = (n - 1) * s ** 2 / mean
    res["D"] = s ** 2 / mean
    res["disp_chi2"] = disp
    res["disp_p"] = float(2 * min(stats.chi2.cdf(disp, n - 1), stats.chi2.sf(disp, n - 1)))

    # So sánh với phép đo 100 s (quy về 10 s)
    N100 = d["N100"]
    res["N100"] = N100
    res["N100_per10"] = N100 * T_SHORT / T_LONG
    res["u_N100_per10"] = math.sqrt(N100) * T_SHORT / T_LONG
    diff = res["N100_per10"] - mean
    res["diff100"] = diff
    res["z100"] = diff / math.hypot(res["u_N100_per10"], res["u_mean"])

    # Tốc độ đếm trung bình tích lũy r_c(j)
    j = np.arange(1, n + 1)
    res["rc"] = (np.cumsum(x) / (j * T_SHORT)).tolist()

    # Bảng tần suất + Poisson/Gauss lý thuyết (chuẩn hóa theo n)
    k = np.arange(x.min(), x.max() + 1)
    obs = np.array([(x == v).sum() for v in k])
    poi = n * stats.poisson.pmf(k, mean)
    gau = n * stats.norm.pdf(k, mean, math.sqrt(mean))
    res["k"], res["obs"] = k.tolist(), obs.tolist()
    res["poi"], res["gau"] = poi.tolist(), gau.tolist()
    res["chi2_poi"], res["bins"] = chi2_eq5(obs, poi)
    res["chi2_gau"], _ = chi2_eq5(obs, gau)
    res["ndf"] = res["bins"] - 1  # 1 tham số (µ) ước lượng từ dữ liệu
    res["p_poi"] = float(stats.chi2.sf(res["chi2_poi"], res["ndf"]))
    res["p_gau"] = float(stats.chi2.sf(res["chi2_gau"], res["ndf"]))
    sd = math.sqrt(mean)
    res["pearson_poi"] = chi2_pearson(x, lambda v: stats.poisson.cdf(v, mean))
    # Gauss rời rạc hóa: bin nguyên k ↔ khoảng (k−0.5, k+0.5]
    res["pearson_gau"] = chi2_pearson(x, lambda v: stats.norm.cdf(v + 0.5, mean, sd))

    # Độ lệch lớn nhất so với trung bình
    dev_hi, dev_lo = x.max() - mean, mean - x.min()
    if dev_hi >= dev_lo:
        p1 = stats.poisson.sf(x.max() - 1, mean)        # P(X ≥ max)
        res["maxdev"] = {"value": int(x.max()), "dev": dev_hi, "side": "trên", "p1": float(p1)}
    else:
        p1 = stats.poisson.cdf(x.min(), mean)           # P(X ≤ min)
        res["maxdev"] = {"value": int(x.min()), "dev": -dev_lo, "side": "dưới", "p1": float(p1)}
    res["maxdev"]["nsig"] = abs(res["maxdev"]["dev"]) / math.sqrt(mean)
    res["maxdev"]["p100"] = float(1 - (1 - res["maxdev"]["p1"]) ** n)
    # Số lần kỳ vọng thấy giá trị ≥ max (hoặc ≤ min) trong 100 lần
    res["maxdev"]["expected"] = n * res["maxdev"]["p1"]

    # Monte Carlo: 1000 lần thử với µ = x̄
    mc = rng.poisson(mean, N_MC)
    res["mc"] = {"mean": float(mc.mean()), "s": float(mc.std(ddof=1)),
                 "u_mean": float(mc.std(ddof=1) / math.sqrt(N_MC)),
                 "min": int(mc.min()), "max": int(mc.max())}
    km = np.arange(mc.min(), mc.max() + 1)
    om = np.array([(mc == v).sum() for v in km])
    res["mc"]["k"], res["mc"]["obs"] = km.tolist(), om.tolist()
    res["mc"]["chi2_poi"], res["mc"]["bins"] = chi2_eq5(om, N_MC * stats.poisson.pmf(km, mean))
    MC_DIR.mkdir(exist_ok=True)
    np.savetxt(MC_DIR / f"mc_{d['key']}.csv", np.column_stack([km, om]), fmt="%d",
               delimiter=",", header="so_dem,tan_suat", comments="")
    np.savetxt(MC_DIR / f"mc_{d['key']}_raw.csv", mc, fmt="%d", header="so_dem", comments="")

    # Phân tán của s khi lặp lại thí nghiệm 100 lần đo (N_REP mẫu MC)
    reps = rng.poisson(mean, (N_REP, n))
    s_reps = reps.std(axis=1, ddof=1)
    m_reps = reps.mean(axis=1)
    res["mc_rep"] = {"s_mean": float(s_reps.mean()), "s_std": float(s_reps.std(ddof=1)),
                     "s_pct": float((s_reps < s).mean() * 100),
                     "m_std": float(m_reps.std(ddof=1)), "s_all": s_reps.tolist()}
    return res, mc


def fig_running(results):
    fig, axes = plt.subplots(3, 1, figsize=(7, 8), sharex=True)
    for ax, (key, r) in zip(axes, results.items()):
        j = np.arange(1, r["n"] + 1)
        rc = np.array(r["rc"])
        ax.plot(j, rc, "o-", ms=2.5, lw=1, color="#1f5f99", label=r"$r_c(j)$")
        ax.axhline(r["rate_cps"], color="#c0392b", ls="--", lw=1,
                   label=fr"$\bar r$ = {r['rate_cps']:.3f} s$^{{-1}}$")
        band = np.sqrt(r["mean"] * j) / (j * T_SHORT)  # ±1σ Poisson của r_c(j)
        ax.fill_between(j, r["rate_cps"] - band, r["rate_cps"] + band, color="#c0392b", alpha=0.12,
                        label=r"±1σ = $\sqrt{\mu j}/(j\,t)$")
        ax.set_ylabel("Tốc độ đếm (s$^{-1}$)")
        ax.set_title(r["label"], fontsize=10)
        ax.legend(fontsize=8, loc="upper right")
        ax.grid(alpha=0.3)
    axes[-1].set_xlabel("Số lần đo j")
    fig.tight_layout()
    fig.savefig(FIG / "fig1_running_mean.png")
    plt.close(fig)


def fig_hist(key, r, idx, with_gauss):
    k = np.array(r["k"])
    obs = np.array(r["obs"])
    kk = np.arange(max(0, k.min() - 3), k.max() + 4)
    fig, ax = plt.subplots(figsize=(7, 4.3))
    ax.plot(kk, r["n"] * stats.poisson.pmf(kk, r["mean"]), "-", color="#127a7a", lw=1.6,
            label=fr"Poisson, µ = {r['mean']:.2f}")
    if with_gauss:
        xs = np.linspace(kk.min(), kk.max(), 400)
        ax.plot(xs, r["n"] * stats.norm.pdf(xs, r["mean"], math.sqrt(r["mean"])), "--",
                color="#d35400", lw=1.4, label=fr"Gauss, µ = σ² = {r['mean']:.2f}")
    m = obs > 0
    ax.errorbar(k[m], obs[m], yerr=np.sqrt(obs[m]), fmt="o", ms=4, color="black", capsize=2,
                elinewidth=1, label="Thực nghiệm (±√n)")
    ax.set_xlabel("Số đếm trong 10 s")
    ax.set_ylabel("Số lần xuất hiện")
    ax.set_title(r["label"], fontsize=10)
    ax.grid(alpha=0.3)
    ax.legend(fontsize=8)
    fig.tight_layout()
    fig.savefig(FIG / f"fig{idx}_hist_{key}.png")
    plt.close(fig)


def fig_mc(results, mcs):
    fig, axes = plt.subplots(1, 3, figsize=(11, 3.6))
    for ax, (key, r) in zip(axes, results.items()):
        mc = mcs[key]
        bins = np.arange(mc.min() - 0.5, mc.max() + 1.5)
        ax.hist(mc, bins=bins, density=True, color="#9ab7d3", edgecolor="#5a7fa5", lw=0.4,
                label=f"Monte Carlo ({N_MC} lần)")
        x = np.array(r["x"])
        bx = np.arange(x.min() - 0.5, x.max() + 1.5)
        ax.hist(x, bins=bx, density=True, histtype="step", color="black", lw=1.3,
                label=f"Thực nghiệm ({r['n']} lần)")
        kk = np.arange(mc.min(), mc.max() + 1)
        ax.plot(kk, stats.poisson.pmf(kk, r["mean"]), "-", color="#c0392b", lw=1.2, label="Poisson")
        ax.set_title(r["label"], fontsize=9)
        ax.set_xlabel("Số đếm trong 10 s")
        ax.grid(alpha=0.3)
    axes[0].set_ylabel("Xác suất")
    axes[0].legend(fontsize=7)
    fig.tight_layout()
    fig.savefig(FIG / "fig5_montecarlo.png")
    plt.close(fig)


def fig_s_spread(results):
    fig, axes = plt.subplots(1, 3, figsize=(11, 3.4))
    for ax, (key, r) in zip(axes, results.items()):
        s_all = np.array(r["mc_rep"]["s_all"])
        ax.hist(s_all, bins=40, color="#c9d9e8", edgecolor="#6b8fb5", lw=0.4)
        ax.axvline(r["s"], color="#c0392b", lw=1.6, label=f"s thực nghiệm = {r['s']:.2f}")
        ax.axvline(r["sqrt_mean"], color="black", ls="--", lw=1, label=f"√µ = {r['sqrt_mean']:.2f}")
        ax.set_title(r["label"], fontsize=9)
        ax.set_xlabel("Độ lệch chuẩn s của mẫu 100 lần đo")
        ax.legend(fontsize=7)
        ax.grid(alpha=0.3)
    axes[0].set_ylabel(f"Số mẫu MC (trên {N_REP})")
    fig.tight_layout()
    fig.savefig(FIG / "fig6_s_spread.png")
    plt.close(fig)


def main():
    FIG.mkdir(exist_ok=True)
    rng = np.random.default_rng(SEED)
    data = load()
    results, mcs = {}, {}
    for key, d in data.items():
        d["key"] = key
        results[key], mcs[key] = analyse(d, rng)
        results[key]["x"] = d["x"].tolist()

    fig_running(results)
    fig_hist("S1", results["S1"], 2, with_gauss=False)
    fig_hist("S2", results["S2"], 3, with_gauss=True)
    fig_hist("S3", results["S3"], 4, with_gauss=True)
    fig_mc(results, mcs)
    fig_s_spread(results)

    out = {k: {kk: vv for kk, vv in v.items() if kk not in ("rc",)} for k, v in results.items()}
    for v in out.values():
        v["mc_rep"] = {kk: vv for kk, vv in v["mc_rep"].items() if kk != "s_all"}
    out["_meta"] = {"T_SHORT": T_SHORT, "T_LONG": T_LONG, "N_MC": N_MC, "N_REP": N_REP, "SEED": SEED}
    (OUT_DIR / "results.json").write_text(json.dumps(out, ensure_ascii=False, indent=1, default=float))

    for k, r in results.items():
        print(f"{k} {r['label']}: x̄={r['mean']:.2f}±{r['u_mean']:.2f}  s={r['s']:.2f}  √x̄={r['sqrt_mean']:.2f}"
              f"  D={r['D']:.3f} (p={r['disp_p']:.3f})  100s→{r['N100_per10']:.1f}±{r['u_N100_per10']:.2f}"
              f" z={r['z100']:.2f}  χ²P={r['chi2_poi']:.2f} χ²G={r['chi2_gau']:.2f} bins={r['bins']}"
              f"  maxdev={r['maxdev']}  MC: x̄={r['mc']['mean']:.2f} s={r['mc']['s']:.2f}"
              f"  s-pct={r['mc_rep']['s_pct']:.1f}%")


if __name__ == "__main__":
    main()
